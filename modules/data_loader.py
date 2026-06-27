# 数据读取模块
# 负责从 JSON / Excel 文件中读取客户数据，并统一转换为 list[dict]。
# 后续入库、质量检查和 AI prompt 构造都基于字典字段访问，因此这里先统一数据形态。

from openpyxl import load_workbook
import json

def load_excel_file(file_name):
    """
    读取 Excel 客户数据，并转换为 list[dict]。

    输入：
        file_name: Excel 文件路径，要求第一行是字段名。

    输出：
        成功时返回客户字典列表；文件不存在时返回 None。
    """
    try:
        workbook = load_workbook(file_name)
        sheet = workbook.active

        # 第一行作为字段名，后续行作为客户记录，这样 Excel 和 JSON 能共用同一套字段。
        rows = list(sheet.iter_rows(values_only=True))
        headers = rows[0]
        customers = []

        for row in rows[1:]:
            # 跳过完全空白的行，避免把空行误导入数据库。
            if all(value is None for value in row):
                continue

            # 将表头和当前行一一配对，得到 {"字段名": "字段值"} 形式的客户字典。
            customer = dict(zip(headers, row))
            customers.append(customer)

        return customers

    except FileNotFoundError:
        print(f"文件不存在,请先创建{file_name}")
        return None

def load_customers(file_name):
    """
    根据文件后缀自动选择读取方式。

    输入：
        file_name: 待导入的客户数据文件路径，当前支持 .json 和 .xlsx。

    输出：
        成功时返回 list[dict]；文件类型不支持或读取失败时返回 None。
    """
    # 入口函数只判断文件类型，具体解析逻辑交给对应函数，方便后续扩展 CSV 等格式。
    if file_name.endswith(".json"):
        return load_json_file(file_name)

    elif file_name.endswith(".xlsx"):
        return load_excel_file(file_name)
        
    else:
        print("暂不支持该文件类型")
        return None

def load_json_file(file_name):
    """
    读取 JSON 客户数据文件。

    输入：
        file_name: JSON 文件路径，内容通常是客户字典列表。

    输出：
        成功时返回 Python 数据结构；文件不存在或 JSON 格式错误时返回 None。
    """
    try:
        with open(
            file_name,
            "r",
            encoding="utf-8"
        ) as file:
            data = json.load(file)
            return data

    except FileNotFoundError:
        print(f"文件不存在,请先创建{file_name}")
        return None

    except json.JSONDecodeError:
        # JSON 格式错误会导致后续无法按字段取值，因此在导入前直接拦截。
        print("JSON格式错误,请检查文件内容")
        return None
